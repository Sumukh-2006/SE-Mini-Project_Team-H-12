#!/usr/bin/env python3
"""Create version 1.1 documents from preserved baseline requirements and real evidence."""
import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
STORIES=json.loads((ROOT/"deliverables/planning/backlog.json").read_text())
SUMMARY=json.loads((ROOT/"deliverables/evidence/verification-summary.json").read_text())
METRICS=json.loads((ROOT/"deliverables/evidence/system-metrics.json").read_text())
if any(c['exit_code'] for c in SUMMARY['checks']):raise SystemExit("Evidence has failures; do not author a passing report")

def new_document(title, purpose):
    doc=Document();section=doc.sections[0]
    section.page_width=Inches(8.27);section.page_height=Inches(11.69)
    section.top_margin=section.bottom_margin=Inches(.7)
    section.left_margin=section.right_margin=Inches(.75)
    normal=doc.styles['Normal'];normal.font.name="Liberation Sans";normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.line_spacing=1.08
    normal.paragraph_format.keep_together=True
    for root in (doc.styles.element,doc.element):
        for border in list(root.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for name,size in [('Title',24),('Heading 1',15),('Heading 2',12)]:
        style=doc.styles[name];style.font.name="Liberation Sans";style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.keep_with_next=True
    doc.add_paragraph(title,'Title')
    doc.add_paragraph('Hospital Management System   Team H12   Version 1.1   12 October 2026')
    doc.add_paragraph(purpose)
    doc.add_paragraph('Prepared locally for team review. The submitted version 1.0 remains unchanged; GitHub publication and team review evidence are pending.')
    doc.core_properties.author="Shashank D"
    doc.core_properties.title=title
    return doc

def heading(doc,text):doc.add_paragraph(text,'Heading 1')
def paragraph(doc,text):doc.add_paragraph(text)
def bullets(doc,items):
    for text in items:doc.add_paragraph(text,'List Bullet')

def table(doc,headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
    if widths:
        for i,width in enumerate(widths):t.columns[i].width=Inches(width)
    for i,text in enumerate(headers):
        t.rows[0].cells[i].text=text
        for run in t.rows[0].cells[i].paragraphs[0].runs:run.bold=True
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells=t.add_row().cells
        for i,value in enumerate(values):cells[i].text=str(value)
    for row in t.rows:
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for i,cell in enumerate(row.cells):
            if widths:cell.width=Inches(widths[i])
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(3)
                for run in p.runs:run.font.size=Pt(9.5)
    return t

def save(doc,folder,name):
    path=ROOT/'docs'/folder/name;path.parent.mkdir(parents=True,exist_ok=True);doc.save(path);print(path)

def srs():
    doc=new_document('Software Requirements Specification','This specification retains all 29 submitted requirements and defines the implementation choices needed to deliver the local console application. It is the acceptance baseline for the two sprint packages.')
    heading(doc,'Product and scope')
    paragraph(doc,'The system manages patient records, doctor schedules and appointments in a self-contained C++17 console application. SQLite persists the data on local disk. Administrators maintain doctors, weekly availability and users and generate reports. Receptionists manage patients and appointments. Doctors search patients, view their own schedules and complete their own consultations. Patients are managed records, not login accounts.')
    bullets(doc,['No external database server or web service is required. Billing, pharmacy and online payments are outside scope.','Use a presentation layer for menus, a business layer for authorization and validation, and a SQLite data access layer.','First-run setup creates one administrator; subsequent accounts require an authenticated administrator. No working database or production credentials are versioned.'])
    heading(doc,'Functional requirements')
    groups=[('AUTH','Authentication and authorization'),('PAT','Patient management'),('DOC','Doctor management'),('APP','Appointment management'),('REP','Reporting')]
    for group,title in groups:
        doc.add_paragraph(title,'Heading 2')
        rows=[(s['id'],s['requirement']) for s in STORIES[:29] if s['id'].startswith('FR-'+group)]
        table(doc,['Requirement ID','Requirement'],rows,[1.25,5.52])
    heading(doc,'Non functional and security requirements')
    table(doc,['Requirement ID','Requirement'],[(s['id'],s['requirement']) for s in STORIES[:29] if not s['id'].startswith('FR-')],[1.25,5.52])
    heading(doc,'Acceptance clarifications')
    bullets(doc,[
        'Patient IDs expose an alphanumeric value such as P000001 while retaining integer database keys. Doctor and appointment IDs are generated integer keys.',
        'Dates use YYYY-MM-DD with calendar validation. Times use HH:MM; appointments last 15 minutes and start on quarter-hour boundaries. Availability has one interval per doctor per weekday, with Sunday as 0.',
        'Only Scheduled appointments may be canceled, rescheduled or completed. Cancellation keeps history and frees its slot; completion retains the slot. Failed rescheduling leaves the original date and time intact.',
        'Doctor logins are uniquely linked to a doctor record. Only that doctor can view their scheduled appointments or complete them. Receptionists and Doctors can search patient records as the SRS permits.',
        'Three invalid logins end the current process attempt session. A new process begins a new attempt session. Logout clears authenticated state.',
        'Admin daily totals use the appointment\'s current scheduled date and status. Total booked includes Scheduled, Canceled and Completed records. CSV and text exports contain patient lists or daily totals.',
        'Passwords are salted SHA-256 using OpenSSL, matching the submitted algorithm. All user SQL values are bound parameters. Sensitive business operations enforce roles before accessing data.',
        'Performance and memory thresholds are measured on the documented Ubuntu WSL workstation. The release gate includes full test coverage, crash rollback, static checks and Valgrind.'])
    heading(doc,'Traceability and change control')
    paragraph(doc,'The complete requirement-to-component-to-test mapping is deliverables/planning/Traceability.md. The local backlog contains these 29 requirements and 10 supporting tasks. Changes from the original design and added completion workflow are recorded in docs/Part2/Deviations.md. The GitHub Project, teammate assignments and contribution history must be published and verified separately.')
    save(doc,'SRS','HMS_SRS_v1.1.docx')

def architecture():
    doc=new_document('Software Architecture and Design','This document describes the implemented three-layer architecture, SQLite schema, service contracts and transactional rules. It adds the missing scheduling and doctor-account relationships while retaining the console-only C++ design.')
    heading(doc,'Layered architecture')
    table(doc,['Layer','Responsibilities and source'],[
        ('Presentation','src/main.cpp prompts, role menus, input conversion, hidden Linux password input and readable errors.'),
        ('Business logic','HMS owns private session state and checks permissions. Auth, Patients, Doctors, Appointments and Reports are implemented in separate source modules.'),
        ('Data access','Database, Statement and Transaction own connections, bound prepared statements and rollback-on-exit transactions. SQLite holds persistent records.')],[1.35,5.42])
    paragraph(doc,'Console UI calls business services; services validate and authorize before preparing data access statements. No HTTP endpoints or remote APIs exist. Services return typed records or throw errors. The UI reports those errors and returns to the menu; end-of-input exits safely.')
    heading(doc,'Database schema')
    table(doc,['Table','Fields and constraints'],[
        ('Doctors','id integer primary key; name, specialization, contact required.'),
        ('Users','id; unique username; password_hash; random salt; role; nullable unique doctor_id. Doctor roles require a linked doctor; other roles forbid a doctor link.'),
        ('Patients','id integer key; unique public_id; name, age 0 to 130, gender, contact and blood_group.'),
        ('Availability','doctor_id and weekday form the key; start_min and end_min hold a valid increasing daily interval.'),
        ('Appointments','id; patient_id and doctor_id foreign keys; appt_date, appt_time and status. Status is Scheduled, Canceled or Completed.')],[1.35,5.42])
    bullets(doc,['Foreign keys are enabled on every connection. WAL and synchronous FULL protect committed writes; missing directories and schema are initialized on startup.','A partial unique index on doctor/date/time applies to all appointments except Canceled. Lookup indexes cover patient name, doctor specialization and appointment date.','Patient registration inserts the integer key and public ID in one transaction. All user-supplied values are bound to prepared statements; schema SQL is constant.'])
    heading(doc,'Service interfaces')
    table(doc,['Subsystem','Operations and authorization'],[
        ('Authentication','bootstrap on an empty Users table; login; logout; provision under Admin. Private session and process attempt counter.'),
        ('Patients','registerPatient and updateContact under Receptionist; searchPatients under Receptionist or Doctor.'),
        ('Doctors','addDoctor and setAvailability under Admin; doctors filter under Admin or Receptionist.'),
        ('Appointments','book, cancel and reschedule under Receptionist; mySchedule and complete under linked Doctor.'),
        ('Reports','dailyReport, exportPatients and exportDaily under Admin. Paths must end in .csv or .txt.')],[1.35,5.42])
    heading(doc,'Appointment booking sequence')
    bullets(doc,['Receptionist submits patient ID, doctor ID, calendar date and quarter-hour time.','Business service checks role and begins an immediate transaction, then resolves the patient and checks weekly hours and active slot occupancy.','The prepared insert creates a Scheduled appointment. The unique index provides an additional conflict guard. Commit succeeds before a success result is returned; any failure rolls back.'])
    heading(doc,'Rescheduling and consultation sequence')
    bullets(doc,['Receptionist identifies a Scheduled appointment and supplies a new date/time.','Within one immediate transaction, the service checks the same doctor\'s availability and conflicts, excluding the appointment being moved.','The service updates date/time and commits. Validation, conflict or database failure leaves the original booking intact.','A linked Doctor may mark their own Scheduled appointment Completed. Canceled or Completed records cannot be rescheduled or completed again.'])
    heading(doc,'Security and lifecycle')
    paragraph(doc,'OpenSSL RAND_bytes generates 16 random salt bytes. The stored hash is SHA-256 of the hexadecimal salt, a separator and the password. Password comparison uses CRYPTO_memcmp. Linux interactive input suppresses password echo. RAII finalizes statements, closes connections and rolls back unfinished transactions. No production password, working database or generated demo secret is included in the delivered repository. Test fixtures contain test-only credentials.')
    paragraph(doc,'Salted SHA-256 satisfies the submitted architecture; a production system should use a slow password KDF and additional account protections. This academic local program uses process-level login denial after three failures, as clarified in the updated SRS.')
    heading(doc,'Compatibility and deviations')
    paragraph(doc,'The verified target is Ubuntu 24.04 under WSL2 with GCC 13.3, SQLite 3.45.1 and OpenSSL 3.0.13. CMake requires C++17 and the SQLite, OpenSSL and GoogleTest development packages. Native Windows and macOS execution have not been validated. Consult docs/Part2/Setup.md and Deviations.md for user workflows and the documented additions to the original schema.')
    save(doc,'Architecture_Design','HMS_ArchDesign_v1.1.docx')

def testplan():
    doc=new_document('Software Test Plan and Results','This test plan retains the 14 submitted test cases, extends coverage to all 29 requirements and reports the executed local results. Hosted CI and team PR review evidence remain separate pending deliverables.')
    heading(doc,'Strategy and environment')
    paragraph(doc,'Use GoogleTest for business and data access behavior, real pseudoterminal console tests for menus and input handling, cppcheck for static analysis, Valgrind for memory errors and definite leaks, and a production DAL helper killed during an uncommitted write for crash recovery. Network load testing is outside scope.')
    paragraph(doc,'Verified environment: Ubuntu 24.04 under WSL2, release C++17 build with GCC 13.3, SQLite 3.45.1, OpenSSL 3.0.13 and GoogleTest 1.14. Run python3 tools/verify_local.py after building. It stores exit codes, raw outputs and GoogleTest XML under deliverables/evidence.')
    heading(doc,'Retained baseline test cases')
    original=Document(ROOT/'docs/TestPlan/HMS_TestPlan_v1.0.docx')
    rows=[]
    for t in original.tables:
        for row in t.rows:
            values=[c.text.strip() for c in row.cells]
            if values and re.fullmatch(r'TC-[A-Z]+-\d{2}',values[0]):rows.append(values)
    assert len(rows)==14
    table(doc,['Test ID and requirement','Scenario and expected result'],[(v[0]+'\n'+v[1],v[2]+'\nExpected: '+v[3]) for v in rows],[1.45,5.32])
    paragraph(doc,'TC-NFR-02 uses the external Valgrind verdict, not a standalone GoogleTest assertion, to verify zero definitely lost memory. TC-NFR-01 uses a 10,000-record fixture; system checks also time actual console search and result display.')
    heading(doc,'Additional scenarios')
    bullets(doc,['Service-side role enforcement and doctor isolation; first-run bootstrap restrictions; salted password storage; literal SQL-like inputs and wildcard-safe name search.','Unique IDs, invalid demographics/contact/blood groups, non-existent records, calendar and slot boundaries, and availability changes that would exclude booked appointments.','Transactional reschedule failure, slot release after cancellation, state transition restrictions, completion ownership and correct daily status counts.','CSV quoting and text content, invalid report filenames, write failures and unauthorized exports; missing database startup and persistence across connections.','Abrupt process termination rolls back an uncommitted write; database integrity remains intact. The ordinary application workflow is checked under Valgrind separately from the test executable.'])
    heading(doc,'Executed local results')
    table(doc,['Check','Observed result'],[
        ('GoogleTest',f"{SUMMARY['tests']} passed; {SUMMARY['failures']} failures; {SUMMARY['errors']} errors."),
        ('Console and static analysis','Real menu workflow, invalid numeric input, three-attempt denial and cppcheck passed.'),
        ('ID query and display',f"10,000-record database: {METRICS['query_wall_seconds']:.4f} seconds; limit 2 seconds."),
        ('Broad name query and display',f"9,999 matching records displayed in {METRICS['broad_query_and_display_seconds']:.4f} seconds; limit 2 seconds."),
        ('Normal application RSS',f"{METRICS['normal_process_rss_kb']} KB; limit 51,200 KB. Measured on the application, outside Valgrind."),
        ('Abrupt termination',METRICS['crash_recovery']),
        ('Memory safety','Valgrind test-suite and actual console workflow: zero errors and zero definitely lost bytes.'),
        ('Hosted and team evidence','Not performed: GitHub Project, hosted CI, PRs, reviews, merges and uploads.')],[1.65,5.12])
    heading(doc,'Complete requirements traceability')
    table(doc,['Requirement ID','Verification'],[(s['id'],s['tests']) for s in STORIES[:29]],[1.25,5.52])
    heading(doc,'Release acceptance and limitations')
    paragraph(doc,'Local acceptance requires all checks passing, no Valgrind errors or definite leaks, both NFR thresholds met, successful crash rollback, preserved baseline documents and verified demo duration. Measurements characterize this workstation, not every supported operating system or concurrent workload. The final team release also requires actual assigned-member PRs, owner reviews, hosted checks and documentation/video publication after authorization.')
    save(doc,'TestPlan','HMS_TestPlan_v1.1.docx')

if __name__=='__main__':srs();architecture();testplan()
